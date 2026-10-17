import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)      # #0F172A
    CARD_BG = RGBColor(30, 41, 59)      # #1E293B
    CARD_BORDER = RGBColor(51, 65, 85)  # #334155
    CYAN = RGBColor(56, 189, 248)       # #38BDF8
    EMERALD = RGBColor(16, 185, 129)    # #10B981
    CORAL = RGBColor(239, 68, 68)       # #EF4444
    AMBER = RGBColor(245, 158, 11)      # #F59E0B
    WHITE = RGBColor(248, 250, 252)     # #F8FAFC
    GRAY = RGBColor(148, 163, 184)      # #94A3B8
    CODE_BG = RGBColor(6, 9, 15)        # #06090F

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.color.rgb = BG_DARK
        return bg

    def add_header(slide, tag_text, title_text, slide_num):
        # Header Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN

        # Header Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(10), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Slide Counter
        counter_box = slide.shapes.add_textbox(Inches(11.0), Inches(0.4), Inches(1.5), Inches(0.35))
        tf_c = counter_box.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = f"Slide {slide_num} / 10"
        p_c.alignment = PP_ALIGN.RIGHT
        p_c.font.size = Pt(11)
        p_c.font.color.rgb = GRAY

        # Divider line
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.02))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.color.rgb = CARD_BORDER

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Big Tag
    tag = s1.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(10), Inches(0.4))
    p = tag.text_frame.paragraphs[0]
    p.text = "⚡ OPERATING SYSTEMS HACKATHON • 17.10.2026"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Main Title
    t1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(1.6))
    tf1 = t1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "Dining Philosophers Deadlock Prevention System"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "A Multi-Threaded POSIX Concurrency, Deadlock Elimination & Starvation Prevention Suite"
    p1_sub.font.size = Pt(17)
    p1_sub.font.color.rgb = GRAY

    # Team Box Container
    add_card(s1, Inches(1.0), Inches(3.6), Inches(11.333), Inches(3.1), bg_color=CARD_BG)
    
    t_head = s1.shapes.add_textbox(Inches(1.3), Inches(3.8), Inches(10), Inches(0.4))
    p_th = t_head.text_frame.paragraphs[0]
    p_th.text = "TEAM: OCTO_OSSP  |  PROJECT PARTICIPATION & ARCHITECTURAL ROLES"
    p_th.font.size = Pt(13)
    p_th.font.bold = True
    p_th.font.color.rgb = CYAN

    members = [
        ("G Jashwanth Reddy", "System Architect & C Concurrency Lead", "POSIX pthreads, mutex management, Dijkstra Resource Ordering hierarchy, Coffman circular wait detection."),
        ("Mohammad Zaid Amaan", "IPC Protocol & Server Lead", "Node.js process supervision, incremental JSON streaming over stdio, full-duplex WebSocket broadcast hub."),
        ("M Pramod Krishna", "Systems Visualization & Analytics Lead", "4-Page SPA dashboard, SVG circular table geometry, hover diagnostics tooltips, and comparative telemetry.")
    ]

    for idx, (name, role, desc) in enumerate(members):
        col_left = Inches(1.3 + idx * 3.65)
        add_card(s1, col_left, Inches(4.3), Inches(3.45), Inches(2.1), bg_color=RGBColor(20, 29, 47), border_color=CARD_BORDER)
        
        mbox = s1.shapes.add_textbox(col_left + Inches(0.15), Inches(4.4), Inches(3.15), Inches(1.9))
        mtf = mbox.text_frame
        mtf.word_wrap = True
        
        p_name = mtf.paragraphs[0]
        p_name.text = name
        p_name.font.size = Pt(15)
        p_name.font.bold = True
        p_name.font.color.rgb = WHITE
        
        p_role = mtf.add_paragraph()
        p_role.text = role
        p_role.font.size = Pt(11)
        p_role.font.bold = True
        p_role.font.color.rgb = CYAN
        
        p_desc = mtf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Problem Statement & Motivation", "The Challenge of Concurrent Synchronization", 2)

    # Problem Statement Callout Banner
    add_card(s2, Inches(0.8), Inches(1.7), Inches(11.733), Inches(1.2), bg_color=RGBColor(24, 34, 53), border_color=CYAN)
    callout = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(1.0))
    ctf = callout.text_frame
    ctf.word_wrap = True
    cp1 = ctf.paragraphs[0]
    cp1.text = 'OFFICIAL PROBLEM STATEMENT:'
    cp1.font.size = Pt(11)
    cp1.font.bold = True
    cp1.font.color.rgb = CYAN
    cp2 = ctf.add_paragraph()
    cp2.text = '"Implement and visualize different synchronization strategies to prevent deadlock and starvation among competing processes/threads."'
    cp2.font.size = Pt(16)
    cp2.font.bold = True
    cp2.font.color.rgb = WHITE

    # Left Column: The Dilemma
    add_card(s2, Inches(0.8), Inches(3.2), Inches(5.7), Inches(3.7))
    lbox = s2.shapes.add_textbox(Inches(1.0), Inches(3.35), Inches(5.3), Inches(3.4))
    ltf = lbox.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "⚠️ The Synchronization Challenge"
    lp1.font.size = Pt(17)
    lp1.font.bold = True
    lp1.font.color.rgb = CORAL

    points_left = [
        "Concurrent threads compete for shared, bounded, non-preemptible resources.",
        "Uncoordinated acquisition causes race conditions and critical section corruption.",
        "When all threads acquire partial resources and wait for others, the system freezes in permanent Deadlock.",
        "Simple deadlock prevention does NOT prevent Starvation: some threads may be indefinitely bypassed by faster peers."
    ]
    for pt in points_left:
        p = ltf.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY

    # Right Column: Our Engineering Solution
    add_card(s2, Inches(6.8), Inches(3.2), Inches(5.733), Inches(3.7))
    rbox = s2.shapes.add_textbox(Inches(7.0), Inches(3.35), Inches(5.3), Inches(3.4))
    rtf = rbox.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = "🎯 Our Engineering Approach"
    rp1.font.size = Pt(17)
    rp1.font.bold = True
    rp1.font.color.rgb = EMERALD

    points_right = [
        "Genuine C Kernel Code: Multi-threaded C with POSIX pthreads on Ubuntu Linux / WSL2 (zero fake frontend simulations).",
        "Deterministic Deadlock Demo: Demonstrates circular wait across all 5 threads using POSIX barrier synchronization.",
        "Deadlock Prevention: Implementation and mathematical proof of Dijkstra's Resource Ordering hierarchy.",
        "Starvation Prevention: Monitor condition variables and FIFO arrival queue guaranteeing bounded wait times."
    ]
    for pt in points_right:
        p = rtf.add_paragraph()
        p.text = "✓ " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 3: THE DINING PHILOSOPHERS PROBLEM
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Theoretical Model", "The Classic Dining Philosophers Problem", 3)

    # 3 Cards Layout
    cards_s3 = [
        ("Classic Model (Dijkstra 1965)", CYAN, [
            "5 Philosophers arranged around a circular dining table.",
            "Each philosopher alternates between THINKING and EATING.",
            "5 Forks situated at the midpoints between adjacent philosophers.",
            "Dining requires BOTH adjacent forks (Left & Right) simultaneously."
        ]),
        ("System Constraints & Invariants", AMBER, [
            "Forks are mutually exclusive resources protected by POSIX mutexes.",
            "No two adjacent philosophers can eat at the same time.",
            "Maximum Concurrent Diners: floor(5 / 2) = 2 philosophers.",
            "Left Fork of P_i is F_i; Right Fork of P_i is F_{(i+1)%5}."
        ]),
        ("Operating Systems Mapping", EMERALD, [
            "Philosopher Threads <==> Concurrent Process / Worker Threads.",
            "Forks <==> Shared Non-Preemptible Hardware / System Resources.",
            "Meals Completed <==> Process Throughput & Task Completion.",
            "Waiting Time <==> Scheduling Latency & Starvation Metric."
        ])
    ]

    for idx, (title, col, bullets) in enumerate(cards_s3):
        c_left = Inches(0.8 + idx * 3.95)
        add_card(s3, c_left, Inches(1.8), Inches(3.8), Inches(5.1))
        
        box = s3.shapes.add_textbox(c_left + Inches(0.2), Inches(2.0), Inches(3.4), Inches(4.7))
        tf = box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        
        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(12)
            pb.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 4: THREE-TIER SYSTEM ARCHITECTURE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "System Design & Architecture", "Three-Tier Decoupled Architecture", 4)

    tiers = [
        ("Component A: C Simulation Engine", "Native Linux / POSIX C Core", CYAN, [
            "Spawns 5 independent POSIX threads (pthread_create).",
            "Manages 5 fork mutexes (pthread_mutex_t) and condition variables.",
            "Implements 3 synchronization engines (Naive, Resource Ordering, Fair FIFO).",
            "Emits high-resolution JSON Line telemetry events over stdout (monotonic clock)."
        ]),
        ("Component B: Communication Server", "Node.js & Process Supervision", AMBER, [
            "Supervises the compiled C binary inside Ubuntu Linux / WSL2.",
            "Parses and validates newline-delimited JSON streams in real time.",
            "Maintains aggregate state snapshots and rolling log buffer.",
            "Full-duplex WebSocket (ws://localhost:3000) relays commands and telemetry."
        ]),
        ("Component C: Interactive Web Dashboard", "HTML5, CSS3 & SVG Vector Engine", EMERALD, [
            "4 Dedicated SPA Navigation Pages (Simulation, Comparison, Logs, Concepts).",
            "Live Circular Dining Table rendered with exact physical fork midpoints.",
            "Multi-layer SVG lines: Topology guides, Acquired emerald, Requested amber.",
            "Rich interactive tooltips, hover diagnostics, and real-time Chart.js analytics."
        ])
    ]

    for idx, (title, sub, col, bullets) in enumerate(tiers):
        t_top = Inches(1.7 + idx * 1.75)
        add_card(s4, Inches(0.8), t_top, Inches(11.733), Inches(1.55))
        
        # Left title box
        tbox = s4.shapes.add_textbox(Inches(1.0), t_top + Inches(0.15), Inches(3.6), Inches(1.2))
        ttf = tbox.text_frame
        ttf.word_wrap = True
        tp1 = ttf.paragraphs[0]
        tp1.text = title
        tp1.font.size = Pt(14)
        tp1.font.bold = True
        tp1.font.color.rgb = col
        tp2 = ttf.add_paragraph()
        tp2.text = sub
        tp2.font.size = Pt(11)
        tp2.font.color.rgb = GRAY
        
        # Right bullets box
        rbox = s4.shapes.add_textbox(Inches(4.8), t_top + Inches(0.1), Inches(7.5), Inches(1.3))
        rtf = rbox.text_frame
        rtf.word_wrap = True
        for b_idx, b in enumerate(bullets):
            bp = rtf.paragraphs[0] if b_idx == 0 else rtf.add_paragraph()
            bp.text = "• " + b
            bp.font.size = Pt(11)
            bp.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 5: COFFMAN CONDITIONS & CONTROLLED DEADLOCK
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "Deadlock Demonstration", "Verification of the Four Coffman Conditions", 5)

    # Left Column: The 4 Conditions
    add_card(s5, Inches(0.8), Inches(1.7), Inches(6.0), Inches(5.2))
    cbox = s5.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.6), Inches(4.8))
    ctf = cbox.text_frame
    ctf.word_wrap = True
    
    cp = ctf.paragraphs[0]
    cp.text = "Coffman Conditions for Deadlock"
    cp.font.size = Pt(17)
    cp.font.bold = True
    cp.font.color.rgb = CORAL

    coffman = [
        ("1. Mutual Exclusion", "Forks cannot be shared; each fork mutex is held exclusively by 1 thread."),
        ("2. Hold and Wait", "Each philosopher holds their left fork while waiting to acquire their right fork."),
        ("3. No Preemption", "Forks cannot be forcibly confiscated from philosophers while eating or waiting."),
        ("4. Circular Wait", "P0 -> F1 -> P1 -> F2 -> P2 -> F3 -> P3 -> F4 -> P4 -> F0 -> P0 forms a closed loop.")
    ]
    for c_title, c_desc in coffman:
        p1 = ctf.add_paragraph()
        p1.text = c_title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = ctf.add_paragraph()
        p2.text = c_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    # Right Column: Controlled Demonstration & Recovery
    add_card(s5, Inches(7.1), Inches(1.7), Inches(5.433), Inches(5.2))
    dbox = s5.shapes.add_textbox(Inches(7.3), Inches(1.9), Inches(5.0), Inches(4.8))
    dtf = dbox.text_frame
    dtf.word_wrap = True

    dp = dtf.paragraphs[0]
    dp.text = "Controlled Deadlock Engineering"
    dp.font.size = Pt(17)
    dp.font.bold = True
    dp.font.color.rgb = AMBER

    d_points = [
        ("Deterministic Synchronization Barrier:", "In educational demo mode, a pthread_barrier_t forces all 5 threads to synchronize and claim left forks simultaneously."),
        ("Dedicated Cycle Watcher Thread:", "A background detection algorithm inspects fork ownership matrices and flags the exact circular wait dependency."),
        ("Safe Non-Blocking Recovery:", "Philosopher threads poll stop/switch flags in 20ms slices using pthread_mutex_trylock, avoiding unrecoverable OS kernel freezing."),
        ("Real-time UI Visualization:", "The web dashboard turns red, draws pulsing circular dependency arrows, and populates the exact cycle breakdown chips.")
    ]
    for d_title, d_desc in d_points:
        p1 = dtf.add_paragraph()
        p1.text = d_title
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = dtf.add_paragraph()
        p2.text = d_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 6: DEADLOCK PREVENTION - RESOURCE ORDERING
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Deadlock Prevention Strategy", "Dijkstra's Resource Ordering Hierarchy", 6)

    # 2 Column layout
    # Left: The Algorithm & Rule
    add_card(s6, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2))
    ro_box = s6.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.8))
    ro_tf = ro_box.text_frame
    ro_tf.word_wrap = True

    p = ro_tf.paragraphs[0]
    p.text = "Total Ordering on Resources"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = CYAN

    ro_text = [
        ("Global Resource Hierarchy:", "Define a strict linear order: F_0 < F_1 < F_2 < F_3 < F_4."),
        ("Acquisition Rule:", "A philosopher must ALWAYS acquire min(Left, Right) before attempting max(Left, Right)."),
        ("Asymmetric Philosopher P4:", "P0 needs F0, F1 -> acquires F0 then F1.\nP1 needs F1, F2 -> acquires F1 then F2.\nP2 needs F2, F3 -> acquires F2 then F3.\nP3 needs F3, F4 -> acquires F3 then F4.\nP4 needs F4, F0 -> min(4, 0) = 0! Acquires F0 FIRST, then F4!"),
        ("Elimination of Circular Wait:", "By forcing P4 to request F0 first, the circular wait condition is broken before it can ever form.")
    ]
    for h, d in ro_text:
        p1 = ro_tf.add_paragraph()
        p1.text = h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = ro_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    # Right: Mathematical Proof
    add_card(s6, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    proof_box = s6.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.8))
    ptf = proof_box.text_frame
    ptf.word_wrap = True

    p = ptf.paragraphs[0]
    p.text = "Mathematical Proof of Deadlock-Freedom"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    proof_points = [
        ("1. Directed Resource Graph:", "Let G = (V, E) be the Resource Allocation Graph where edges represent thread wait dependencies."),
        ("2. Monotonic Index Invariant:", "Every edge (P_i -> F_k) in a wait dependency must satisfy index(F_held) < index(F_requested)."),
        ("3. Acyclicity:", "A cycle requires a sequence F_a < F_b < ... < F_n < F_a, which implies F_a < F_a. This is a mathematical contradiction!"),
        ("4. Theorem:", "Because the Resource Allocation Graph is guaranteed to be a Directed Acyclic Graph (DAG), Circular Wait is IMPOSSIBLE. The system is provably Deadlock-Free!")
    ]
    for h, d in proof_points:
        p1 = ptf.add_paragraph()
        p1.text = h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = ptf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 7: STARVATION PREVENTION - FAIR FIFO MONITOR
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Starvation Prevention Strategy", "Fair FIFO Monitor with Condition Variables", 7)

    add_card(s7, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2))
    fifo_l = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.8))
    fl_tf = fifo_l.text_frame
    fl_tf.word_wrap = True

    p = fl_tf.paragraphs[0]
    p.text = "Why Deadlock Prevention != Fairness"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = AMBER

    fl_points = [
        ("The Starvation Vulnerability:", "Under Resource Ordering, deadlock is eliminated, but threads can still suffer STARVATION due to OS scheduling bias."),
        ("Conspiracy Starvation:", "Philosophers P0 and P2 can alternate eating, keeping Fork 1 perpetually locked. Philosopher P1 waits indefinitely and starves."),
        ("Need for Admission Control:", "Eliminating starvation requires a centralized monitor that tracks request arrival times and enforces fair scheduling.")
    ]
    for h, d in fl_points:
        p1 = fl_tf.add_paragraph()
        p1.text = h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = fl_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    add_card(s7, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    fifo_r = s7.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.8))
    fr_tf = fifo_r.text_frame
    fr_tf.word_wrap = True

    p = fr_tf.paragraphs[0]
    p.text = "The Fair FIFO Monitor Solution"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    fr_points = [
        ("Arrival FIFO Queue:", "Hungry philosophers are registered in a strict first-in-first-out arrival queue."),
        ("Condition Variable Synchronization:", "Each philosopher thread blocks on a condition variable (pthread_cond_t can_eat_cond[i])."),
        ("Admission Invariant:", "A thread is admitted to eat IF AND ONLY IF:\n1. Both adjacent forks are free, AND\n2. No earlier waiting thread in the queue shares either fork."),
        ("Starvation-Free Proof:", "Since every hungry philosopher moves forward in the finite queue as preceding meals finish (t_eat <= T_max), maximum wait time is strictly bounded!")
    ]
    for h, d in fr_points:
        p1 = fr_tf.add_paragraph()
        p1.text = h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = fr_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 8: DASHBOARD & USER INTERFACE DESIGN
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Visualization & User Interface", "Interactive 4-Page System Dashboard", 8)

    ui_cards = [
        ("Page 1: Live Simulation", CYAN, [
            "Circular Dining Table with 5 Philosophers & 5 Forks.",
            "Accurate geometry: each fork at exact midpoint of adjacent philosophers.",
            "Multi-layer SVG lines: Topology guide (dotted), Acquired (emerald), Requested (amber), Deadlock (pulsing red).",
            "Real-time telemetry indicators: Meals, Wait Time, Engine Uptime."
        ]),
        ("Page 2: Algorithm Comparison", EMERALD, [
            "Side-by-side benchmark of Resource Ordering vs Fair FIFO vs Naive.",
            "Live meal distribution bar chart comparing thread throughput.",
            "Variance & Fairness indicators measuring starvation resistance.",
            "Comparative summary table with average and max wait times."
        ]),
        ("Page 3: Event Log & Monitoring", AMBER, [
            "Millisecond-accurate chronological event telemetry stream.",
            "Filter pills: Lifecycle, Mutex Locks, State Changes, Deadlocks.",
            "Per-philosopher search queries and autoscroll toggle.",
            "Log export to JSON file for offline academic analysis."
        ]),
        ("Page 4: OS Concepts & Architecture", RGBColor(168, 85, 247), [
            "Interactive textbook on POSIX synchronization primitives.",
            "Detailed explanations of Dijkstra's algorithm & Coffman conditions.",
            "Three-tier architecture data flow documentation.",
            "Technical guide designed for computer science evaluations."
        ])
    ]

    for idx, (title, col, bullets) in enumerate(ui_cards):
        col_idx = idx % 2
        row_idx = idx // 2
        
        c_left = Inches(0.8 + col_idx * 5.95)
        c_top = Inches(1.7 + row_idx * 2.65)
        
        add_card(s8, c_left, c_top, Inches(5.75), Inches(2.45))
        
        box = s8.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.15), Inches(5.35), Inches(2.15))
        tf = box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col
        
        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 9: TESTING, VERIFICATION & BENCHMARKS
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Verification & Quality Assurance", "Comprehensive Automated Testing Suite", 9)

    add_card(s9, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2))
    test_l = s9.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(5.3), Inches(4.8))
    tl_tf = test_l.text_frame
    tl_tf.word_wrap = True

    p = tl_tf.paragraphs[0]
    p.text = "C Engine Unit Tests: 46 / 46 Passed"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    c_tests = [
        ("Circular Topology Invariants (10 Tests):", "Explicitly verifies P0..P4 adjacent fork pairs and exclusive sharing of F0..F4."),
        ("Fork Mutual Exclusion:", "Proves that trylock fails when another thread holds the fork."),
        ("Resource Ordering Progress:", "Validates concurrent dining across threads without deadlock halts."),
        ("Controlled Deadlock Detection:", "Confirms circular wait cycle detection and graceful termination."),
        ("Fair FIFO Fairness Verification:", "Empirically tests that meal variance remains bounded (spread <= 1).")
    ]
    for h, d in c_tests:
        p1 = tl_tf.add_paragraph()
        p1.text = "✓ " + h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = tl_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = GRAY

    add_card(s9, Inches(6.8), Inches(1.7), Inches(5.733), Inches(5.2))
    test_r = s9.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.3), Inches(4.8))
    tr_tf = test_r.text_frame
    tr_tf.word_wrap = True

    p = tr_tf.paragraphs[0]
    p.text = "E2E Integration Tests: 8 / 8 Passed"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN

    e2e_tests = [
        ("Child Process Supervision:", "Spawns, monitors, and stops C process in WSL/Linux seamlessly."),
        ("Command IPC Protocol:", "Validates START, PAUSE, RESUME, STOP, and RESET commands."),
        ("Runtime Topology Assertions:", "Asserts that every active fork ownership strictly obeys adjacency."),
        ("Zero Memory Leaks & Orphans:", "Full pthread_join cleanup guarantees no orphan threads or zombie processes.")
    ]
    for h, d in e2e_tests:
        p1 = tr_tf.add_paragraph()
        p1.text = "✓ " + h
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p2 = tr_tf.add_paragraph()
        p2.text = d
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = GRAY

    # =========================================================================
    # SLIDE 10: CONCLUSION & SUMMARY
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)
    add_header(s10, "Summary & Conclusion", "Key Takeaways & Hackathon Deliverables", 10)

    # 3 Summary Cards
    conclusions = [
        ("Academic & Practical Impact", CYAN, [
            "Bridges theoretical operating systems concepts with production-quality POSIX C code.",
            "Visualizes abstract synchronization primitives (mutexes, condition variables) in real time.",
            "Provides an intuitive educational tool for understanding deadlocks and race conditions."
        ]),
        ("Technical Excellence", EMERALD, [
            "Three genuine synchronization algorithms implemented and verified.",
            "Decoupled, scalable three-tier architecture (C Engine + Node.js + Web SPA).",
            "100% automated test coverage (46 C unit tests, 8 integration tests).",
            "Responsible resource management with zero orphan threads or memory leaks."
        ]),
        ("Ready for Live Demonstration", AMBER, [
            "Live Web Dashboard at http://localhost:3000.",
            "Complete GitHub Repository with comprehensive setup guides and documentation.",
            "Timed 5-7 minute live demo script covering all evaluation rubrics.",
            "Thank you! We are ready for your questions and live demonstration."
        ])
    ]

    for idx, (title, col, bullets) in enumerate(conclusions):
        c_left = Inches(0.8 + idx * 3.95)
        add_card(s10, c_left, Inches(1.8), Inches(3.8), Inches(5.1))
        
        box = s10.shapes.add_textbox(c_left + Inches(0.2), Inches(2.0), Inches(3.4), Inches(4.7))
        tf = box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col
        
        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = "• " + b
            pb.font.size = Pt(12)
            pb.font.color.rgb = GRAY

    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == '__main__':
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.abspath(os.path.join(script_dir, '..'))
    
    # Save in docs/presentation/, docs/, and project root
    out1 = os.path.join(project_root, 'docs', 'presentation', 'OCTO_OSSP_Dining_Philosophers_Presentation.pptx')
    out2 = os.path.join(project_root, 'docs', 'OCTO_OSSP_Dining_Philosophers_Presentation.pptx')
    out3 = os.path.join(project_root, 'OCTO_OSSP_Dining_Philosophers_Presentation.pptx')
    
    create_deck(out1)
    create_deck(out2)
    create_deck(out3)
